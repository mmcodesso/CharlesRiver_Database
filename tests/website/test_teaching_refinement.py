"""Website teaching claims against the unchanged, checksum-matched release.

These independent equivalents do not execute native Excel or generator fixtures.
Run with the repository .venv and a timeout of at least 3600 seconds.
"""

from __future__ import annotations

from collections import defaultdict
import csv
from datetime import date, datetime, time
from decimal import Decimal
import hashlib
import json
import math
from pathlib import Path
import re
import sqlite3

from openpyxl import load_workbook
import pytest
import yaml


ROOT = Path(__file__).resolve().parents[2]
SQLITE = ROOT / "outputs/CharlesRiver.sqlite"
RELEASE_SHA256 = "de861968e07412d71032b6ed6a0fb8eb9825cf6911886b1597c4a628a914474c"
REGISTERS = [
    json.loads((ROOT / f"planning/website-redesign/{area}/verified-examples.json").read_text())
    for area in ("tutorials", "cases")
]
EXAMPLES = [example for register in REGISTERS for example in register["examples"]]
REPORTS = yaml.safe_load((ROOT / "config/report_catalog.yaml").read_text())["reports"]
APPROVED_SCOPE = """
    SELECT EmployeeID, PayrollPeriodID, WorkCenterID, RegularHours, OvertimeHours
    FROM TimeClockEntry
    WHERE PayrollPeriodID IN (1,2)
      AND WorkDate BETWEEN '2024-01-01' AND '2024-01-28'
      AND ClockStatus='Approved'
"""


def cents(value: float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"))


@pytest.fixture(scope="module")
def release_db():
    assert SQLITE.is_file(), "Install the published v2027.1 SQLite; never generate a replacement"
    with SQLITE.open("rb") as source:
        assert hashlib.file_digest(source, "sha256").hexdigest() == RELEASE_SHA256
    assert all(register["sha256"] == RELEASE_SHA256 for register in REGISTERS)
    db = sqlite3.connect(SQLITE.resolve().as_uri() + "?mode=ro", uri=True)
    db.row_factory = sqlite3.Row
    db.execute("PRAGMA query_only=ON")
    yield db
    db.close()
    with SQLITE.open("rb") as source:
        assert hashlib.file_digest(source, "sha256").hexdigest() == RELEASE_SHA256


def test_worked_inventory_preserves_all_105_published_queries():
    assert [len(register["examples"]) for register in REGISTERS] == [43, 62]
    assert len({example["name"] for example in EXAMPLES}) == 105
    published = {}
    for page in {example["page"] for example in EXAMPLES}:
        text = (ROOT / page).read_text()
        for name, sql in re.findall(r'```sql title="([^"]+)"\n([\s\S]*?)\n```', text):
            assert name not in published, name
            published[name] = (page, sql)
    assert len(published) == 105
    for example in EXAMPLES:
        assert published[example["name"]] == (example["page"], example["sql"])


@pytest.mark.parametrize("example", EXAMPLES, ids=lambda example: example["name"])
def test_existing_worked_result_reexecutes_read_only(release_db, example):
    assert re.match(r"^(SELECT|WITH)\b", example["sql"], re.I)
    assert not re.search(r"\b(INSERT|UPDATE|DELETE|DROP|ATTACH|CREATE|REPLACE)\b", example["sql"], re.I)
    cursor = release_db.execute(example["sql"])
    actual = [dict(row) for row in cursor]
    expected = example["expected"]
    if expected:
        assert [column[0] for column in cursor.description] == list(expected[0])
    assert actual == expected


def test_f01_null_centers_preserve_every_employee_period_and_pay(release_db):
    detail = [dict(row) for row in release_db.execute(APPROVED_SCOPE)]
    assert len(detail) == 782
    groups = defaultdict(list)
    for row in detail:
        groups[row["EmployeeID"], row["PayrollPeriodID"]].append(row)
    assert len(groups) == 79

    # Model separate Power Query references; filtering is confined to the count branch.
    identified = {
        key: len({row["WorkCenterID"] for row in rows if row["WorkCenterID"] is not None})
        for key, rows in groups.items()
    }
    missing = {key: sum(row["WorkCenterID"] is None for row in rows) for key, rows in groups.items()}
    affected = {key for key, count in missing.items() if count}
    assert len(affected) == 15
    assert {(109, 1), (109, 2)} <= affected
    assert all(identified[key] == 0 for key in affected)
    assert all(len({row["WorkCenterID"] for row in groups[key]}) == identified[key] + 1 for key in affected)

    sql_groups = {
        (row["EmployeeID"], row["PayrollPeriodID"]): row
        for row in release_db.execute(f"""
            SELECT EmployeeID, PayrollPeriodID,
                   COUNT(DISTINCT WorkCenterID) AS Centers,
                   SUM(WorkCenterID IS NULL) AS Unassigned,
                   SUM(RegularHours) AS Regular, SUM(OvertimeHours) AS Overtime
            FROM ({APPROVED_SCOPE}) GROUP BY EmployeeID,PayrollPeriodID
        """)
    }
    assert set(sql_groups) == set(groups)
    for key, rows in groups.items():
        assert identified[key] == sql_groups[key]["Centers"]
        assert missing[key] == sql_groups[key]["Unassigned"]
        assert cents(sum(row["RegularHours"] for row in rows)) == cents(sql_groups[key]["Regular"])
        assert cents(sum(row["OvertimeHours"] for row in rows)) == cents(sql_groups[key]["Overtime"])

    payroll = [dict(row) for row in release_db.execute("SELECT * FROM PayrollRegister WHERE PayrollPeriodID IN (1,2)")]
    bridge = [dict(row, Centers=identified.get((row["EmployeeID"], row["PayrollPeriodID"]))) for row in payroll]
    assert len(bridge) == len(payroll) == 209
    assert len({row["PayrollRegisterID"] for row in bridge}) == 209
    assert sum(row["Centers"] is None for row in bridge) == 130
    assert {(row["EmployeeID"], row["PayrollPeriodID"]) for row in bridge if row["Centers"] == 0} == affected
    for period, expected_net in ((1, "104320.59"), (2, "220069.69")):
        for field in ("GrossPay", "EmployeeWithholdings", "NetPay", "EmployerPayrollTax", "EmployerBenefits"):
            assert cents(sum(row[field] for row in bridge if row["PayrollPeriodID"] == period)) == cents(
                sum(row[field] for row in payroll if row["PayrollPeriodID"] == period)
            )
        assert cents(sum(row["NetPay"] for row in bridge if row["PayrollPeriodID"] == period)) == Decimal(expected_net)


def test_f02_audit_counts_entries_and_distinct_people_separately(release_db):
    rows = release_db.execute("""
        SELECT t.TimeClockEntryID,t.EmployeeID,a.OvertimeApprovalID AS LinkedApproval
        FROM TimeClockEntry t LEFT JOIN OvertimeApproval a USING(OvertimeApprovalID)
        WHERE t.WorkDate>='2026-10-01' AND t.WorkDate<'2027-01-01'
          AND t.ClockStatus='Approved' AND t.OvertimeHours>0
    """).fetchall()
    missing = [row for row in rows if row["LinkedApproval"] is None]
    assert len(rows) == len({row["TimeClockEntryID"] for row in rows}) == 1895
    assert len(missing) == 336
    assert len({row["EmployeeID"] for row in missing}) == 71
    assert round(100 * len(missing) / len(rows), 2) == 17.73


def test_f03_calendar_match_and_capacity_have_different_grain_checks(release_db):
    center_only = release_db.execute("""
        SELECT COUNT(*) FROM WorkOrderOperationSchedule s
        JOIN WorkCenterCalendar c ON c.WorkCenterID=s.WorkCenterID
        WHERE s.WorkOrderOperationScheduleID=5
    """).fetchone()[0]
    matched = release_db.execute("""
        SELECT s.WorkCenterID,s.ScheduleDate,s.ScheduledHours,c.AvailableHours
        FROM WorkOrderOperationSchedule s JOIN WorkCenterCalendar c
          ON c.WorkCenterID=s.WorkCenterID AND c.CalendarDate=s.ScheduleDate
        WHERE s.WorkOrderOperationScheduleID=5
    """).fetchall()
    assert center_only == 1180
    assert len(matched) == 1
    assert tuple(matched[0])[:3] == (4, "2024-01-19", 0.42)

    daily = release_db.execute("""
        WITH scheduled AS (
          SELECT WorkCenterID,ScheduleDate,COUNT(*) AS Operations,SUM(ScheduledHours) AS Scheduled
          FROM WorkOrderOperationSchedule GROUP BY WorkCenterID,ScheduleDate
        ) SELECT s.Operations,s.Scheduled,c.AvailableHours
        FROM scheduled s JOIN WorkCenterCalendar c
          ON c.WorkCenterID=s.WorkCenterID AND c.CalendarDate=s.ScheduleDate
        WHERE s.WorkCenterID=1 AND s.ScheduleDate='2024-01-22'
    """).fetchall()
    assert len(daily) == 1
    assert tuple(daily[0]) == (4, 6.45, 101.76)
    repeated = release_db.execute("""
        SELECT SUM(c.AvailableHours) FROM WorkOrderOperationSchedule s JOIN WorkCenterCalendar c
          ON c.WorkCenterID=s.WorkCenterID AND c.CalendarDate=s.ScheduleDate
        WHERE s.WorkCenterID=1 AND s.ScheduleDate='2024-01-22'
    """).fetchone()[0]
    assert cents(repeated) == Decimal("407.04")
    assert cents(repeated) == cents(4 * daily[0]["AvailableHours"])
    assert release_db.execute("""
        SELECT COUNT(*) FROM (SELECT WorkCenterID,CalendarDate FROM WorkCenterCalendar
        GROUP BY WorkCenterID,CalendarDate HAVING COUNT(*)<>1)
    """).fetchone()[0] == 0


def test_f04_complete_close_balances_and_reverses_each_pl_balance(release_db):
    rows = release_db.execute("""
        SELECT a.AccountID,a.AccountNumber,a.AccountType,g.Debit,g.Credit
        FROM GLEntry g JOIN Account a USING(AccountID)
        WHERE g.SourceDocumentType='JournalEntry' AND g.SourceDocumentID=846
    """).fetchall()
    assert len(rows) == 42
    assert cents(sum(row["Debit"] for row in rows)) == Decimal("27503887.96")
    assert cents(sum(row["Credit"] for row in rows)) == Decimal("27503887.96")
    selected = {row["AccountNumber"]: row for row in rows}
    for number, debit, credit in (
        (4010, "15673061.54", "0.00"),
        (4060, "0.00", "145655.35"),
        (6070, "0.00", "268131.51"),
        (8010, "0.00", "5110308.27"),
    ):
        assert cents(selected[number]["Debit"]) == Decimal(debit)
        assert cents(selected[number]["Credit"]) == Decimal(credit)
    assert cents(sum(selected[number]["Debit"] - selected[number]["Credit"] for number in (4010,4060,6070,8010))) != 0
    for row in rows:
        if row["AccountType"] not in {"Revenue", "Expense"}:
            continue
        prior_balance = release_db.execute("""
            SELECT COALESCE(SUM(Debit-Credit),0) FROM GLEntry
            WHERE AccountID=? AND PostingDate<='2024-12-31'
              AND NOT(SourceDocumentType='JournalEntry' AND SourceDocumentID=846)
        """, (row["AccountID"],)).fetchone()[0]
        assert cents(prior_balance + row["Debit"] - row["Credit"]) == 0
        assert row["Credit"] > 0 if prior_balance > 0 else row["Debit"] > 0


def test_f05_payroll_summary_aggregates_before_ledger_comparison(release_db):
    registers = release_db.execute("SELECT * FROM PayrollRegister WHERE PayrollPeriodID=1 AND CostCenterID=1").fetchall()
    assert len(registers) == 3
    assert cents(sum(row["NetPay"] for row in registers)) == Decimal("23489.30")
    ledger = release_db.execute("""
        SELECT g.*,a.AccountNumber FROM GLEntry g JOIN Account a USING(AccountID)
        WHERE g.VoucherNumber='PRS-0001-001'
    """).fetchall()
    assert len(ledger) == 6
    assert {(row["VoucherType"], row["SourceDocumentType"], row["SourceDocumentID"], row["SourceLineID"], row["CostCenterID"])
            for row in ledger} == {("PayrollSummary", "PayrollSummary", 1001001, None, 1)}
    assert 1001001 == 1_000_000 + 1 * 1000 + 1
    payable = [row for row in ledger if row["AccountNumber"] == 2030]
    assert len(payable) == 1
    assert cents(payable[0]["Credit"]) == cents(sum(row["NetPay"] for row in registers))
    assert cents(sum(row["Debit"] - row["Credit"] for row in ledger)) == 0
    assert release_db.execute("SELECT COUNT(*) FROM PayrollRegister WHERE PayrollRegisterID=1001001").fetchone()[0] == 0
    # The anomaly-enriched release intentionally retains three balanced payment
    # vouchers whose supporting payment records are missing; classify, do not erase.
    orphan_payments = release_db.execute("""
        SELECT g.SourceDocumentID,COUNT(*) AS Lines,ROUND(SUM(g.Debit-g.Credit),2) AS Balance
        FROM GLEntry g LEFT JOIN PayrollPayment p
          ON p.PayrollPaymentID=g.SourceDocumentID
        WHERE g.SourceDocumentType='PayrollPayment' AND p.PayrollPaymentID IS NULL
        GROUP BY g.SourceDocumentID ORDER BY g.SourceDocumentID
    """).fetchall()
    assert [tuple(row) for row in orphan_payments] == [(1,2,0),(2646,2,0),(5420,2,0)]
    assert release_db.execute("SELECT COUNT(*) FROM GLEntry WHERE SourceDocumentType='PayrollPayment'").fetchone()[0] == 16394
    assert release_db.execute("SELECT COUNT(*) FROM GLEntry WHERE SourceDocumentType='PayrollPaymentSummary'").fetchone()[0] == 0


@pytest.mark.parametrize("start,expected_count", [("2024-03-01",162),("2024-01-01",486)])
def test_f06_date_intervals_include_first_day_exclude_april(release_db, start, expected_count):
    count, first, last = release_db.execute("""
        SELECT COUNT(*),MIN(OrderDate),MAX(OrderDate) FROM SalesOrder
        WHERE OrderDate>=? AND OrderDate<'2024-04-01'
    """, (start,)).fetchone()
    assert count == expected_count
    assert first == start
    assert last == "2024-03-31"
    assert release_db.execute("SELECT COUNT(*) FROM SalesOrder WHERE OrderDate='2024-04-01'").fetchone()[0] > 0
    for date, included in ((start,1),("2024-03-31",1),("2024-04-01",0)):
        assert release_db.execute("SELECT ? >= ? AND ? < '2024-04-01'", (date,start,date)).fetchone()[0] == included


def test_f09_completion_and_close_aggregate_independently_at_work_order(release_db):
    assert "ProductionCompletionLineID" not in {row[1] for row in release_db.execute("PRAGMA table_info(WorkOrderClose)")}
    completion = release_db.execute("""
        SELECT COUNT(*) AS Lines,SUM(l.ExtendedStandardTotalCost) AS Completed
        FROM ProductionCompletion h JOIN ProductionCompletionLine l USING(ProductionCompletionID)
        WHERE h.WorkOrderID=1
    """).fetchone()
    close = release_db.execute("SELECT COUNT(*) AS Closes,SUM(TotalVarianceAmount) AS Variance FROM WorkOrderClose WHERE WorkOrderID=1").fetchone()
    assert completion["Lines"] == 2
    assert close["Closes"] == 1
    assert cents(completion["Completed"]) == Decimal("1019.85")
    assert cents(close["Variance"]) == Decimal("743.16")
    naive = release_db.execute("""
        SELECT SUM(c.TotalVarianceAmount) FROM ProductionCompletion h
        JOIN ProductionCompletionLine l USING(ProductionCompletionID)
        JOIN WorkOrderClose c USING(WorkOrderID) WHERE h.WorkOrderID=1
    """).fetchone()[0]
    assert cents(naive) == Decimal("1486.32")
    correct = release_db.execute("""
        WITH completed AS (
          SELECT h.WorkOrderID,SUM(l.ExtendedStandardTotalCost) AS Completed
          FROM ProductionCompletion h JOIN ProductionCompletionLine l USING(ProductionCompletionID)
          GROUP BY h.WorkOrderID
        ), closed AS (
          SELECT WorkOrderID,SUM(TotalVarianceAmount) AS Variance FROM WorkOrderClose GROUP BY WorkOrderID
        ) SELECT w.WorkOrderID,p.Completed,c.Variance FROM WorkOrder w
          LEFT JOIN completed p USING(WorkOrderID) LEFT JOIN closed c USING(WorkOrderID)
          WHERE w.WorkOrderID=1
    """).fetchall()
    assert len(correct) == 1
    assert cents(correct[0]["Completed"]) == cents(completion["Completed"])
    assert cents(correct[0]["Variance"]) == cents(close["Variance"])


def test_f08_f10_release_inventory_is_explicit(release_db):
    assert release_db.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'").fetchone()[0] == 82
    assert release_db.execute("SELECT COUNT(*) FROM TimeClockPunch").fetchone()[0] == 165137


def test_e01_complete_event_vouchers_and_cutoff_preserve_accounting_scope(release_db):
    expected = {
        ("Shipment",87): ("2024-01-31",6,Decimal("6089.34")),
        ("SalesInvoice",202): ("2024-01-31",5,Decimal("12248.33")),
        ("CashReceipt",173): ("2024-02-22",2,Decimal("952.07")),
        ("CashReceiptApplication",225): ("2024-02-23",2,Decimal("952.07")),
    }
    all_rows = []
    for (kind, identifier), (date, count, debit) in expected.items():
        rows = release_db.execute("""
            SELECT a.AccountNumber,g.PostingDate,g.Debit,g.Credit
            FROM GLEntry g JOIN Account a USING(AccountID)
            WHERE g.SourceDocumentType=? AND g.SourceDocumentID=?
        """, (kind,identifier)).fetchall()
        assert len(rows) == count
        assert {row["PostingDate"] for row in rows} == {date}
        assert cents(sum(row["Debit"] for row in rows)) == debit
        assert cents(sum(row["Credit"] for row in rows)) == debit
        if kind == "CashReceiptApplication":
            assert {row["AccountNumber"] for row in rows} == {1020,2060}
        all_rows.extend(rows)
    assert len(all_rows) == 15
    # A different application already exists at January 31. The selected later
    # receipt/application pair cannot establish the entire invoice's open amount.
    for date, expected_applied in (("2024-01-31", "238.16"),("2024-02-29", "1190.23"),("2024-04-16", "12248.33")):
        amount = release_db.execute("""
            SELECT COALESCE(SUM(AppliedAmount),0) FROM CashReceiptApplication
            WHERE SalesInvoiceID=202 AND ApplicationDate<=?
        """, (date,)).fetchone()[0]
        assert cents(amount) == Decimal(expected_applied)


@pytest.mark.parametrize("page,expected_counts", [
    ("docs/reference/posting.md", [1,1]),
    ("docs/processes/manufacturing.md", [1,1,1]),
    ("docs/learn-and-practice/manufacturing-cost-flow.md", [11,1,3]),
])
def test_new_reference_queries_execute_in_addition_to_original_105(release_db, page, expected_counts):
    snippets = re.findall(r"```sql\n([\s\S]*?)\n```", (ROOT / page).read_text())
    assert len(snippets) == len(expected_counts)
    results = [list(release_db.execute(sql)) for sql in snippets]
    assert [len(rows) for rows in results] == expected_counts
    if page.endswith("manufacturing-cost-flow.md"):
        ledger, variance, periods = results
        # Each selected complete voucher balances before the source rows are
        # summarized for presentation; the query output is at account/voucher grain.
        totals = defaultdict(lambda: [0.0,0.0])
        for row in ledger:
            totals[row["VoucherNumber"]][0] += row["Debit"]
            totals[row["VoucherNumber"]][1] += row["Credit"]
        assert len(totals) == 4
        assert all(cents(debit) == cents(credit) for debit,credit in totals.values())
        row = variance[0]
        assert cents(row["RecordedDirectLabor"]) == Decimal("191.89")
        assert cents(row["StandardLabor"]) == Decimal("185.67")
        assert cents(row["StandardOverhead"]) == Decimal("104.82")
        assert cents(row["LaborDifference"]) == cents(row["DirectLaborVarianceAmount"]) == Decimal("6.22")
        assert cents(row["ConversionVarianceAmount"]) == Decimal("718.23")
        expected = [
            ("January 2024",0,20875,23327.60,-2452.60),
            ("February 1",-2452.60,0,12894.16,-15346.76),
            ("February 2",-15346.76,169382.48,8968.64,145067.08),
        ]
        for row, values in zip(periods,expected):
            assert row["Period"] == values[0]
            assert [cents(row[key]) for key in ("Opening","PeriodDebit","PeriodCredit","Closing")] == [cents(value) for value in values[1:]]
            assert cents(row["Opening"]+row["PeriodDebit"]-row["PeriodCredit"]) == cents(row["Closing"])


def assert_export_value(actual, expected, label):
    # CSV and workbook serialization represent SQL NULL as a blank. Excel also
    # represents an empty text cell as a blank; that format limitation is explicit.
    if expected is None or expected == "":
        assert actual in (None, ""), label
    elif isinstance(actual, (date,datetime)) and isinstance(expected,str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}",expected):
        # XLSX stores typed calendar dates; SQLite/CSV/preview serialize ISO text.
        if isinstance(actual,datetime):
            assert actual.time() == time.min, label
            actual = actual.date()
        assert actual.isoformat() == expected, label
    elif isinstance(expected, (int,float)):
        assert actual not in (None,""), label
        try:
            numeric = float(actual)
        except (ValueError,TypeError):
            pytest.fail(f"{label}: expected numeric {expected!r}; got {actual!r}")
        assert math.isfinite(numeric), label
        # This allows binary/Excel decimal representation noise, far below a cent
        # for the report amounts; it does not round away missing or duplicated rows.
        assert math.isclose(numeric, expected, rel_tol=1e-14, abs_tol=1e-8), f"{label}: {actual!r} != {expected!r}"
    else:
        assert actual == expected, f"{label}: {actual!r} != {expected!r}"


@pytest.mark.parametrize("report", REPORTS, ids=lambda report: report["slug"])
def test_restored_report_binaries_match_exact_canonical_sql(release_db, report):
    directory = ROOT / "static/reports" / report["area"] / report["process_group"] / report["slug"]
    paths = [directory / "preview.json", directory / f'{report["slug"]}.csv', directory / f'{report["slug"]}.xlsx']
    before = {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths}
    sql = (ROOT / report["query_path"]).read_text()
    cursor = release_db.execute(sql)
    columns = [column[0] for column in cursor.description]
    expected = [tuple(row) for row in cursor]
    preview = json.loads(paths[0].read_text())
    assert preview["slug"] == report["slug"]
    assert preview["columns"] == columns
    assert preview["rowCount"] == len(expected)
    assert preview["previewRowLimit"] == report["preview_row_limit"]
    assert preview["previewRowCount"] == len(preview["rows"]) == min(len(expected),report["preview_row_limit"])
    for row_index, (actual_row, expected_row) in enumerate(zip(preview["rows"], expected)):
        assert list(actual_row) == columns
        for column, expected_value in zip(columns,expected_row):
            assert_export_value(actual_row[column],expected_value,f'{report["slug"]}: preview row{row_index} {column}')

    with paths[1].open(encoding="utf-8-sig",newline="") as stream:
        reader = csv.reader(stream)
        assert next(reader) == columns
        csv_rows = list(reader)
    assert len(csv_rows) == len(expected)
    for row_index,(actual_row,expected_row) in enumerate(zip(csv_rows,expected)):
        assert len(actual_row) == len(columns)
        for column,actual_value,expected_value in zip(columns,actual_row,expected_row):
            assert_export_value(actual_value,expected_value,f'{report["slug"]}: CSV row{row_index} {column}')

    workbook = load_workbook(paths[2],read_only=True,data_only=False)
    try:
        assert len(workbook.sheetnames) == 1
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        assert list(next(rows)) == columns
        workbook_rows = list(rows)
        assert len(workbook_rows) == len(expected)
        for row_index,(actual_row,expected_row) in enumerate(zip(workbook_rows,expected)):
            assert len(actual_row) == len(columns)
            for column,actual_value,expected_value in zip(columns,actual_row,expected_row):
                assert_export_value(actual_value,expected_value,f'{report["slug"]}: XLSX row{row_index} {column}')
    finally:
        workbook.close()
    assert {path: hashlib.sha256(path.read_bytes()).hexdigest() for path in paths} == before
